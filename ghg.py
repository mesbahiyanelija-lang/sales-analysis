from docx import Document
from docx.shared import Inches

doc = Document()

doc.add_heading("1. Sales by Analysis Report", 0)

doc.add_heading("1. Sales by Category", level=1)
doc.add_picture(r"C:\Users\mouood system\Desktop\char1.png", width=Inches(5))
doc.add_paragraph("Technology leads sales with $827,455, followed by Furniture (728,658) ans Ofiice Supplies (705,422). Technology has approximately 13% higher sales than Furniture. It is recommended to maintain focus on Technology while monitoring growth in other categories")

doc.add_heading("2. Sales by Region", level=1)
doc.add_picture(r"C:\Users\mouood system\Desktop\char2.png", width=Inches(5))
doc.add_paragraph("West and East regions show the strongest performance. Central and South require furthur investigation to determine whether reduced investment or market expansion is more appropriate.")

doc.add_heading("3. Sales by Region and Category", level=1)
doc.add_picture(r"C:\Users\mouood system\Desktop\char3.png", width=Inches(5))
doc.add_paragraph("technology significantly outperforms other categories in the East region. All three categories perform consitently in West. Central and South regions show lower sales across all categories and require strategic review.")

doc.add_heading("4. Monthly Sales Trend", level=1)
doc.add_picture(r"C:\Users\mouood system\Desktop\char4.png", width=Inches(5))
doc.add_paragraph("Overal sales show an upward trend from 2015 to 2018. A significant peak occurs at the end of each year. However, sales drop sharply  at the beginning of each year, wich requires investigation - wether due to low demand or insufficient investment in the first half of the year.")

doc.add_heading("Conclusion", level=1)
doc.add_paragraph("The Store shows healthy overall growth. Key recommendation: maintain Technology focus, investigate Central and South regions underperformance, and develope strategies to boost first-half yearly sales. ")

doc.save(r"C:\Users\mouood system\Desktop\Sales_Report.docx")
print("Report saved on Desktop!")
