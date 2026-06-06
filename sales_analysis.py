import pandas as pd
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches

# ==================== Load Data ====================
df = pd.read_csv(r"C:\Users\mouood system\Downloads\superstore_final_dataset (1).csv\superstore_final_dataset (1).csv", encoding="latin1")
df["Order_Date"] = pd.to_datetime(df["Order_Date"], dayfirst=True)
df["Month"] = df["Order_Date"].dt.to_period("M")

# ==================== Chart 1: Sales by Category ====================
sales_by_category = df.groupby("Category")["Sales"].sum()
sales_by_category.plot(kind="bar", color="steelblue")
plt.title("Sales by Category")
plt.ylabel("Total Sales")
plt.xlabel("Category")
plt.tight_layout()
plt.savefig(r"C:\Users\mouood system\Desktop\chart1.png")
plt.close()

# ==================== Chart 2: Sales by Region ====================
sales_by_region = df.groupby("Region")["Sales"].sum().sort_values(ascending=False)
sales_by_region.plot(kind="bar", color="steelblue")
plt.title("Sales by Region")
plt.ylabel("Total Sales")
plt.xlabel("Region")
plt.tight_layout()
plt.savefig(r"C:\Users\mouood system\Desktop\chart2.png")
plt.close()

# ==================== Chart 3: Sales by Region and Category ====================
pivot = df.pivot_table(values="Sales", index="Region", columns="Category", aggfunc="sum")
pivot.plot(kind="bar")
plt.title("Sales by Region and Category")
plt.ylabel("Total Sales")
plt.xlabel("Region")
plt.tight_layout()
plt.savefig(r"C:\Users\mouood system\Desktop\chart3.png")
plt.close()

# ==================== Chart 4: Monthly Sales Trend ====================
monthly_sales = df.groupby("Month")["Sales"].sum()
monthly_sales.plot(kind="line", color="steelblue")
plt.title("Monthly Sales Trend")
plt.ylabel("Total Sales")
plt.xlabel("Month")
plt.tight_layout()
plt.savefig(r"C:\Users\mouood system\Desktop\chart4.png")
plt.close()

# ==================== Generate Report ====================
doc = Document()
doc.add_heading("Sales Analysis Report", 0)

doc.add_heading("1. Sales by Category", level=1)
doc.add_picture(r"C:\Users\mouood system\Desktop\chart1.png", width=Inches(5))
doc.add_paragraph("Technology leads sales with $827,455, followed by Furniture ($728,658) and Office Supplies ($705,422). Technology has approximately 13% higher sales than Furniture. It is recommended to maintain focus on Technology while monitoring growth opportunities in other categories.")

doc.add_heading("2. Sales by Region", level=1)
doc.add_picture(r"C:\Users\mouood system\Desktop\chart2.png", width=Inches(5))
doc.add_paragraph("West and East regions show the strongest performance. Central and South regions require further investigation to determine whether reduced investment or market expansion is more appropriate.")

doc.add_heading("3. Sales by Region and Category", level=1)
doc.add_picture(r"C:\Users\mouood system\Desktop\chart3.png", width=Inches(5))
doc.add_paragraph("Technology significantly outperforms other categories in the East region. All three categories perform consistently in West. Central and South regions show lower sales across all categories and require strategic review.")

doc.add_heading("4. Monthly Sales Trend", level=1)
doc.add_picture(r"C:\Users\mouood system\Desktop\chart4.png", width=Inches(5))
doc.add_paragraph("Overall sales show an upward trend from 2015 to 2018. A significant peak occurs at the end of each year. However, sales drop sharply at the beginning of each year, which requires investigation.")

doc.add_heading("Conclusion", level=1)
doc.add_paragraph("The store shows healthy overall growth. Key recommendations: maintain Technology focus, investigate South and Central underperformance, and develop strategies to boost first-half yearly sales.")

doc.save(r"C:\Users\mouood system\Desktop\Sales_Report.docx")
print("All done! Report saved on Desktop!")